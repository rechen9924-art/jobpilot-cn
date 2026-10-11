#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从传入的简历 JSON（默认 schema/example-cv.json）生成干净、可编辑的 Word 简历（.docx）。
单栏专业排版，中文采用微软雅黑，无任何工具署名/水印。
用法: python scripts/build-cv-docx.py <output.docx>
"""
import json
import sys
import os

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "微软雅黑"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)   # 深蓝灰，用于标题/分隔线
DARK = RGBColor(0x22, 0x22, 0x22)


def set_cjk(run, font=FONT):
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), font)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)


def add_bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1F3A5F")
    pBdr.append(bottom)


def set_margins(doc, cm=1.8):
    for section in doc.sections:
        section.top_margin = Cm(cm)
        section.bottom_margin = Cm(cm)
        section.left_margin = Cm(cm)
        section.right_margin = Cm(cm)


def add_runs(p, text, size=10.5, bold=False, color=DARK, italic=False, bg=None):
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    set_cjk(run)
    if bg:
        rPr = run._element.get_or_add_rPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), bg)
        rPr.append(shd)
    return run


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    add_runs(p, text, size=12.5, bold=True, color=ACCENT)
    add_bottom_border(p)
    return p


def bullet_item(doc, text, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(2)
    add_runs(p, "•  ", size=size, color=ACCENT)
    add_runs(p, text, size=size)
    return p


def build(cv, out_path):
    # 兼容 basics 嵌套结构：把 basics 内的字段提升为顶层（不覆盖已有顶层字段）
    basics = cv.get("basics") if isinstance(cv.get("basics"), dict) else {}
    for k, v in basics.items():
        cv.setdefault(k, v)

    doc = Document()
    set_margins(doc, 1.8)

    # 默认正文字体
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # ---- 头部 ----
    name_p = doc.add_paragraph()
    name_p.paragraph_format.space_after = Pt(2)
    add_runs(name_p, cv.get("name", ""), size=20, bold=True, color=ACCENT)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(2)
    add_runs(sub_p, cv.get("title", ""), size=11, bold=True, color=DARK)

    links = cv.get("links") or []
    contact = "　".join(
        x for x in [
            (f"电话：{cv.get('phone','')}" if cv.get("phone") else ""),
            (f"邮箱：{cv.get('email','')}" if cv.get("email") else ""),
            (f"所在地：{cv.get('location','')}" if cv.get("location") else ""),
            *[f"链接：{l}" for l in links],
        ]
        if x
    )
    c_p = doc.add_paragraph()
    c_p.paragraph_format.space_after = Pt(2)
    add_runs(c_p, contact, size=9.5, color=RGBColor(0x55, 0x55, 0x55))

    # ---- 个人简介 ----
    if cv.get("summary"):
        section_heading(doc, "个人简介")
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        add_runs(sp, cv["summary"], size=10.5)

    # ---- 核心能力 ----
    if cv.get("core_capabilities"):
        section_heading(doc, "核心能力")
        cp = doc.add_paragraph()
        cp.paragraph_format.space_after = Pt(6)
        for i, cap in enumerate(cv["core_capabilities"]):
            if i > 0:
                add_runs(cp, "  ", size=10)
            add_runs(cp, cap, size=10, color=RGBColor(0x1D, 0x4E, 0xD8), bg="EEF6FF")

    # ---- 工作经历 ----
    if cv.get("experience"):
        section_heading(doc, "工作经历")
        for e in cv["experience"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            add_runs(p, f"{e.get('role','')} · {e.get('company','')}", size=11, bold=True, color=ACCENT)
            meta = "　".join(
                x for x in [
                    e.get("dates", ""),
                    e.get("location", ""),
                ]
                if x
            )
            if meta:
                mp = doc.add_paragraph()
                mp.paragraph_format.space_after = Pt(2)
                add_runs(mp, meta, size=9.5, color=RGBColor(0x55, 0x55, 0x55))
            for b in e.get("bullets", []):
                bullet_item(doc, b)

    # ---- 项目经历 ----
    if cv.get("projects"):
        section_heading(doc, "项目经历")
        for pr in cv["projects"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            add_runs(p, pr.get("name", ""), size=11, bold=True, color=ACCENT)
            if pr.get("tech"):
                tp = doc.add_paragraph()
                tp.paragraph_format.space_after = Pt(1)
                tp.paragraph_format.left_indent = Cm(0.2)
                add_runs(tp, f"技术栈：{pr['tech']}", size=9.5, color=RGBColor(0x55, 0x55, 0x55), italic=True)
            if pr.get("description"):
                dp = doc.add_paragraph()
                dp.paragraph_format.left_indent = Cm(0.2)
                dp.paragraph_format.space_after = Pt(3)
                add_runs(dp, pr["description"], size=10.5)

    # ---- 教育背景 ----
    if cv.get("education"):
        section_heading(doc, "教育背景")
        for ed in cv["education"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            org = ed.get("org", "")
            title = ed.get("title", "")
            year = ed.get("year", "")
            line = ""
            if org:
                line += org
            if title:
                line += f"　{title}"
            if year:
                line += f"　（{year}）"
            add_runs(p, line, size=10.5, bold=True, color=DARK)
            if ed.get("note"):
                np = doc.add_paragraph()
                np.paragraph_format.left_indent = Cm(0.6)
                np.paragraph_format.space_after = Pt(2)
                add_runs(np, ed["note"], size=10)

    # ---- 专业技能 ----
    if cv.get("skills"):
        section_heading(doc, "专业技能")
        for sk in cv["skills"]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.2)
            p.paragraph_format.space_after = Pt(2)
            cat = sk.get("category", "")
            items = sk.get("items", [])
            add_runs(p, f"【{cat}】 ", size=10.5, bold=True, color=ACCENT)
            add_runs(p, "　·　".join(items), size=10.5)

    doc.save(out_path)
    print("✅ Word 简历已生成:", out_path)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    root = os.path.dirname(here)
    # 用法：python scripts/build-cv-docx.py [cv.json] [output.docx]
    # 默认从 schema/example-cv.json 输出到 output/example-resume.docx
    if len(sys.argv) >= 3:
        cv_path = sys.argv[1]
        out_path = sys.argv[2]
    elif len(sys.argv) == 2:
        # 若参数带 .json 则视为 cv，否则视为 out
        if sys.argv[1].lower().endswith(".json"):
            cv_path = sys.argv[1]
            out_path = os.path.join(root, "output", "example-resume.docx")
        else:
            cv_path = os.path.join(root, "schema", "example-cv.json")
            out_path = sys.argv[1]
    else:
        cv_path = os.path.join(root, "schema", "example-cv.json")
        out_path = os.path.join(root, "output", "example-resume.docx")
    with open(cv_path, "r", encoding="utf-8") as f:
        cv = json.load(f)
    build(cv, out_path)
